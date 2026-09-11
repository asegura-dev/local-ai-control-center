"""Tests for the converter port and its PDF and `.docx` implementations."""

from __future__ import annotations

import zipfile
from collections.abc import Callable
from pathlib import Path

import docx
import pytest
from conftest import pdf_with_streams
from pypdf import PdfReader

from local_ai_control_center.adapters.documents import (
    PAGE_MARKER,
    DocxConverter,
    HiddenText,
    PdfConverter,
    converter_for,
    furniture_in,
    hidden_text_in,
    supported_suffixes,
    without_furniture,
)
from local_ai_control_center.ports.converter import ConversionError, Converter


def test_converter_cannot_be_instantiated() -> None:
    """The abstract contract cannot be used directly."""
    with pytest.raises(TypeError):
        Converter()  # type: ignore[abstract]


def test_converter_for_matches_by_suffix(tmp_path: Path) -> None:
    """A source is matched to a converter by what it is, not by the caller choosing."""
    assert isinstance(converter_for(tmp_path / "paper.pdf"), PdfConverter)
    assert isinstance(converter_for(tmp_path / "chapter.docx"), DocxConverter)


def test_converter_for_ignores_suffix_case(tmp_path: Path) -> None:
    """A shouted extension is the same extension."""
    assert isinstance(converter_for(tmp_path / "PAPER.PDF"), PdfConverter)


def test_converter_for_names_what_is_supported(tmp_path: Path) -> None:
    """An unsupported format fails clearly instead of passing bytes through as text."""
    with pytest.raises(ConversionError) as excinfo:
        converter_for(tmp_path / "notes.rtf")
    message = str(excinfo.value)
    assert "notes.rtf" in message
    assert ".pdf" in message and ".docx" in message


def test_supported_suffixes_covers_every_converter() -> None:
    """The list offered in messages is derived from the converters, not written twice."""
    assert supported_suffixes() == frozenset({".pdf", ".docx"})


def test_pdf_text_is_extracted(make_pdf: Callable[..., Path]) -> None:
    """The words in the document come out of it."""
    path = make_pdf("Attention is all you need")
    assert "Attention is all you need" in converter_for(path).extract_text(path)


def test_pdf_marks_page_boundaries(make_pdf: Callable[..., Path]) -> None:
    """Pages are recovered structure: a thesis has to cite them."""
    path = make_pdf("First page", "Second page")
    text = converter_for(path).extract_text(path)
    assert PAGE_MARKER.format(number=1) in text
    assert PAGE_MARKER.format(number=2) in text
    assert text.index("First page") < text.index("Second page")


def test_pdf_page_marker_cannot_be_read_as_a_prompt_fence(make_pdf: Callable[..., Path]) -> None:
    """The marker is an HTML comment, so it can never look like a document fence."""
    path = make_pdf("Some text")
    assert PAGE_MARKER.format(number=1).startswith("<!--")
    assert "<<<" not in converter_for(path).extract_text(path)


def test_pdf_without_a_text_layer_is_reported(make_pdf: Callable[..., Path]) -> None:
    """A scan is named as a scan, and nothing is returned to be written as an empty file."""
    path = make_pdf(" ")
    with pytest.raises(ConversionError) as excinfo:
        converter_for(path).extract_text(path)
    message = str(excinfo.value)
    assert "no extractable text" in message
    assert "OCR" in message


def test_unreadable_pdf_fails_clearly(tmp_path: Path) -> None:
    """A file that is not a PDF fails with a message, not a library traceback."""
    path = tmp_path / "broken.pdf"
    path.write_text("this is not a PDF at all", encoding="utf-8")
    with pytest.raises(ConversionError) as excinfo:
        converter_for(path).extract_text(path)
    assert "broken.pdf" in str(excinfo.value)


def test_missing_pdf_fails_clearly(tmp_path: Path) -> None:
    """A path that is not there is reported, not raised as a bare OSError."""
    with pytest.raises(ConversionError) as excinfo:
        PdfConverter().extract_text(tmp_path / "absent.pdf")
    assert "absent.pdf" in str(excinfo.value)


def test_docx_text_is_extracted(make_docx: Callable[..., Path]) -> None:
    """Body text comes through as body text."""
    path = make_docx((None, "The chapter reviews prior work."))
    assert "The chapter reviews prior work." in converter_for(path).extract_text(path)


def test_docx_heading_levels_come_from_styles(make_docx: Callable[..., Path]) -> None:
    """Headings are recovered from Word's own styles, never guessed from appearance."""
    path = make_docx(
        ("Heading 1", "Theory"),
        ("Heading 2", "Prior work"),
        (None, "Plain paragraph."),
    )
    text = converter_for(path).extract_text(path)
    assert "# Theory" in text
    assert "## Prior work" in text
    assert "\nPlain paragraph." in text


def test_docx_tables_are_flattened_not_dressed_as_markdown(tmp_path: Path) -> None:
    """Cells are real, so they are kept - as rows, not as a table LACC cannot guarantee."""
    path = tmp_path / "tabla.docx"
    document = docx.Document()
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Method"
    table.rows[0].cells[1].text = "Error"
    table.rows[1].cells[0].text = "Newton"
    table.rows[1].cells[1].text = "1e-8"
    document.save(str(path))

    text = converter_for(path).extract_text(path)
    assert "Method | Error" in text
    assert "Newton | 1e-8" in text
    assert "---" not in text


def test_empty_docx_is_reported(tmp_path: Path) -> None:
    """An empty document is a failure to report, not an empty file to write."""
    path = tmp_path / "vacio.docx"
    docx.Document().save(str(path))
    with pytest.raises(ConversionError) as excinfo:
        converter_for(path).extract_text(path)
    assert "no extractable text" in str(excinfo.value)


def test_docx_that_is_not_a_docx_fails_clearly(tmp_path: Path) -> None:
    """The older `.doc` format is not this format, and the message says so."""
    path = tmp_path / "viejo.docx"
    path.write_bytes(b"not a zip container at all")
    with pytest.raises(ConversionError) as excinfo:
        converter_for(path).extract_text(path)
    assert ".doc" in str(excinfo.value)


def test_docx_does_not_resolve_external_entities(
    make_docx: Callable[..., Path], tmp_path: Path
) -> None:
    """A `.docx` is untrusted XML: an external entity must not be able to read a file.

    The classic XXE. python-docx does not resolve entities today; this test is here so
    that a change of parser, or of library, cannot quietly reintroduce it.
    """
    secret = tmp_path / "secret.txt"
    secret.write_text("PRIVATE-CONTENT-12345", encoding="utf-8")
    base = make_docx((None, "harmless"), name="base.docx")

    hostile = tmp_path / "xxe.docx"
    document_xml = (
        '<?xml version="1.0"?>\n'
        f'<!DOCTYPE r [<!ENTITY xxe SYSTEM "{secret.as_uri()}">]>\n'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        "<w:body><w:p><w:r><w:t>&xxe;</w:t></w:r></w:p></w:body></w:document>"
    )
    with zipfile.ZipFile(base) as source, zipfile.ZipFile(hostile, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "word/document.xml":
                data = document_xml.encode("utf-8")
            target.writestr(item, data)

    try:
        text = DocxConverter().extract_text(hostile)
    except ConversionError:
        return  # refusing to parse it at all is also a correct outcome
    assert "PRIVATE-CONTENT" not in text


def test_a_running_header_carrying_a_page_number_is_recognised() -> None:
    """It is never literally identical, which is why counting repeated lines finds nothing.

    The page number is glued to the front, so `3413European Journal...` and
    `3417European Journal...` are different strings and the same furniture (ADR-036).
    """
    pages = [
        "3413European Journal of Nuclear Medicine\nbody text of the first page",
        "3414European Journal of Nuclear Medicine\nbody text of the second page",
        "3415European Journal of Nuclear Medicine\nbody text of the third page",
    ]
    assert furniture_in(pages) == frozenset({"#European Journal of Nuclear Medicine"})


def test_a_line_on_one_page_of_several_is_content() -> None:
    """Body text does not repeat, and a heuristic that dropped it would be worse than the
    problem it solves."""
    pages = ["unique to the first page", "ordinary text", "more ordinary text"]
    assert furniture_in(pages) == frozenset()


def test_a_single_page_has_no_furniture_to_find() -> None:
    """With one page there is nothing to compare against, so nothing is dropped."""
    assert furniture_in(["a header\nand some body text"]) == frozenset()


def test_furniture_removal_reports_how_many_lines_went() -> None:
    """Ingestion edits rather than only transcribing now, so what it removed is counted.

    A heuristic that quietly deletes text from a document the user keeps is the wrong
    shape for this project (ADR-036).
    """
    pages = ["1 3\nreal text here", "1 3\nmore real text", "1 3\nfurther real text"]
    cleaned, dropped = without_furniture(pages)
    assert dropped == 3
    assert all("1 3" not in page for page in cleaned)
    assert "real text here" in cleaned[0]


def _hidden_in(tmp_path: Path, stream: str) -> tuple[HiddenText, ...]:
    """Build a one-page PDF with that content stream and report what is hidden in it."""
    path = tmp_path / "hidden.pdf"
    path.write_bytes(pdf_with_streams([stream]))
    found, _fragments = hidden_text_in(PdfReader(path))
    return found


def test_an_ordinary_page_hides_nothing(tmp_path: Path) -> None:
    """A detector that fires on normal documents is one people switch off."""
    assert _hidden_in(tmp_path, "BT /F1 12 Tf 20 200 Td (Perfectly ordinary text) Tj ET") == ()


def test_text_drawn_in_an_invisible_mode_is_reported(tmp_path: Path) -> None:
    """Mode 3 draws nothing. It is also what a scan's OCR layer uses, which is why this is
    reported and never judged (ADR-040)."""
    found = _hidden_in(tmp_path, "BT /F1 12 Tf 20 200 Td (Visible) Tj 3 Tr 20 190 Td (HIDE) Tj ET")
    assert any("invisible rendering mode" in item.reason for item in found)


def test_text_too_small_to_read_is_reported_with_its_words(tmp_path: Path) -> None:
    """The words matter: they are what the person needs to see, since the model saw them."""
    found = _hidden_in(
        tmp_path, "BT /F1 12 Tf 20 200 Td (Visible) Tj /F1 0.5 Tf 20 190 Td (HIDDEN TEXT) Tj ET"
    )
    assert any(item.reason == "too small to read" and "HIDDEN TEXT" in item.text for item in found)


def test_text_positioned_off_the_page_is_reported(tmp_path: Path) -> None:
    """Outside the MediaBox is outside what anyone saw, and inside what the model read."""
    found = _hidden_in(tmp_path, "BT /F1 12 Tf 20 200 Td (Visible) Tj 20 -5000 Td (OFF PAGE) Tj ET")
    assert any(item.reason == "positioned off the page" for item in found)


def test_hidden_text_still_reaches_the_extracted_document(tmp_path: Path) -> None:
    """The reason any of this matters, stated as a test.

    Hidden text ends up in what the model is given, and a quotation of it passes the
    grounding check because it genuinely is in the document (ADR-038, ADR-040).
    """
    path = tmp_path / "hidden.pdf"
    path.write_bytes(
        pdf_with_streams(["BT /F1 12 Tf 20 200 Td (Visible) Tj 3 Tr 20 190 Td (PLANTED) Tj ET"])
    )
    assert "PLANTED" in PdfConverter().extract_text(path)


def test_the_proportion_is_reported_alongside_what_was_found(tmp_path: Path) -> None:
    """The denominator is what separates a layout from an attack.

    A document where most fragments sit off the page is telling you about its typesetting;
    one where two of five hundred are hidden is telling you something else, and a bare count
    cannot distinguish them (ADR-040).
    """
    path = tmp_path / "hidden.pdf"
    path.write_bytes(
        pdf_with_streams(
            ["BT /F1 12 Tf 20 200 Td (One) Tj 20 -30 Td (Two) Tj /F1 0.5 Tf 20 -30 Td (Hid) Tj ET"]
        )
    )
    found, fragments = hidden_text_in(PdfReader(path))
    assert len(found) == 1
    assert fragments == 3


def test_the_rendered_size_is_judged_not_the_declared_one(tmp_path: Path) -> None:
    """A real paper set `Tf 1.0` and scaled by 17 in the matrix.

    Reading the declaration alone marked all 123 of its fragments invisible - the false
    positive that gets a detector switched off within a week (ADR-040).
    """
    path = tmp_path / "scaled.pdf"
    path.write_bytes(
        pdf_with_streams(["BT /F1 1 Tf 12 0 0 12 20 200 Tm (Normal sized after scaling) Tj ET"])
    )
    found, _ = hidden_text_in(PdfReader(path))
    assert found == ()
